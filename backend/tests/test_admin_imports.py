import io
import uuid
from decimal import Decimal

import httpx
import openpyxl
import pytest
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from app.auth.models import User
from app.catalog.models import Product, ProductVariant
from app.core.security import create_access_token
from app.imports.sources.xlsx import HEADERS
from app.workers.tasks import apply_import_job
from tests.helpers import s3_public_url_reachable, s3_reachable

pytestmark = pytest.mark.skipif(
    not s3_reachable(), reason="S3/MinIO endpoint is not reachable in this environment"
)


def _slug(prefix: str) -> str:
    return f"{prefix}-{uuid.uuid4().hex[:10]}"


def _build_xlsx(rows: list[list[object]]) -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


async def _make_admin(db_session: AsyncSession) -> User:
    user = User(email=f"admin-{uuid.uuid4().hex[:10]}@example.com", role="admin")
    db_session.add(user)
    await db_session.commit()
    return user


def _admin_headers(user: User) -> dict[str, str]:
    return {"Authorization": f"Bearer {create_access_token(user.id, user.role)}"}


async def _apply_bound_to_test_session(
    db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch, *, import_job_id: str
) -> None:
    # apply_import_job (the worker task) opens its own session via
    # async_session_factory, as a real arq worker would -- bind it to this
    # test's connection so it sees the not-yet-committed-to-the-real-DB setup.
    monkeypatch.setattr(
        "app.workers.tasks.async_session_factory",
        async_sessionmaker(
            bind=db_session.bind, expire_on_commit=False, join_transaction_mode="create_savepoint"
        ),
    )
    await apply_import_job({}, import_job_id=import_job_id)


@pytest.mark.asyncio
async def test_customer_forbidden_from_imports(
    client: httpx.AsyncClient, db_session: AsyncSession
) -> None:
    customer = User(email=f"customer-{uuid.uuid4().hex[:10]}@example.com", role="customer")
    db_session.add(customer)
    await db_session.commit()

    response = await client.get(
        "/v1/admin/imports/template", headers=_admin_headers(customer)
    )

    assert response.status_code == 403


@pytest.mark.asyncio
async def test_template_download_has_expected_headers(
    client: httpx.AsyncClient, db_session: AsyncSession
) -> None:
    admin = await _make_admin(db_session)

    response = await client.get("/v1/admin/imports/template", headers=_admin_headers(admin))

    assert response.status_code == 200
    workbook = openpyxl.load_workbook(io.BytesIO(response.content))
    sheet = workbook.active
    assert sheet is not None
    header_row = next(sheet.iter_rows(max_row=1, values_only=True))
    assert list(header_row) == HEADERS


@pytest.mark.asyncio
async def test_upload_preview_reports_counts_and_row_errors(
    client: httpx.AsyncClient, db_session: AsyncSession
) -> None:
    admin = await _make_admin(db_session)
    category = _slug("Категория")
    sku_ok = _slug("sku")
    name = _slug("Товар")

    file_bytes = _build_xlsx(
        [
            [None, name, category, "Описание", sku_ok, "", "150.00", "", "10", "", ""],
            [None, "Без SKU", category, "Описание", "", "", "150.00", "", "10", "", ""],
            [
                None, "Плохая цена", category, "Описание", _slug("sku"),
                "", "не число", "", "10", "", "",
            ],
        ]
    )

    response = await client.post(
        "/v1/admin/imports/xlsx",
        headers=_admin_headers(admin),
        files={"file": ("test.xlsx", file_bytes, "application/octet-stream")},
    )

    assert response.status_code == 201
    body = response.json()
    preview = body["preview"]
    assert preview["total_rows"] == 3
    assert preview["created_estimate"] == 1
    assert preview["updated_estimate"] == 0
    assert preview["error_count"] == 2
    assert len(preview["rows"]) == 3


@pytest.mark.asyncio
async def test_import_200_rows_creates_products_and_variants(
    client: httpx.AsyncClient, db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    admin = await _make_admin(db_session)
    category = _slug("Массовая категория")
    external_ids = [_slug("EXT") for _ in range(50)]

    rows: list[list[object]] = []
    for external_id in external_ids:
        name = f"Товар {external_id}"
        for variant_index in range(4):
            rows.append(
                [
                    external_id,
                    name,
                    category,
                    "Описание",
                    _slug("sku"),
                    f"вариант={variant_index}",
                    "199.00",
                    "",
                    "10",
                    "",
                    "",
                ]
            )
    assert len(rows) == 200

    upload_response = await client.post(
        "/v1/admin/imports/xlsx",
        headers=_admin_headers(admin),
        files={"file": ("bulk.xlsx", _build_xlsx(rows), "application/octet-stream")},
    )
    assert upload_response.status_code == 201
    import_id = upload_response.json()["import_id"]
    assert upload_response.json()["preview"]["total_rows"] == 200
    assert upload_response.json()["preview"]["created_estimate"] == 200

    apply_response = await client.post(
        f"/v1/admin/imports/{import_id}/apply", headers=_admin_headers(admin)
    )
    assert apply_response.status_code == 200

    await _apply_bound_to_test_session(db_session, monkeypatch, import_job_id=import_id)

    status_response = await client.get(
        f"/v1/admin/imports/{import_id}", headers=_admin_headers(admin)
    )
    body = status_response.json()
    assert body["status"] == "completed"
    assert body["created_count"] == 200
    assert body["error_count"] == 0

    products_count = await db_session.scalar(
        select(func.count()).select_from(Product).where(Product.name.like("Товар EXT-%"))
    )
    assert products_count == 50


@pytest.mark.asyncio
async def test_reimport_with_changed_prices_updates_not_duplicates(
    client: httpx.AsyncClient, db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    admin = await _make_admin(db_session)
    external_id = _slug("EXT")
    name = f"Товар {external_id}"
    category = _slug("Категория")
    sku_a = _slug("sku-a")
    sku_b = _slug("sku-b")

    first_file = _build_xlsx(
        [
            [
                external_id, name, category, "Описание", sku_a,
                "цвет=красный", "300.00", "", "20", "", "",
            ],
            [
                external_id, name, category, "Описание", sku_b,
                "цвет=синий", "350.00", "", "15", "", "",
            ],
        ]
    )

    upload_1 = await client.post(
        "/v1/admin/imports/xlsx",
        headers=_admin_headers(admin),
        files={"file": ("first.xlsx", first_file, "application/octet-stream")},
    )
    import_id_1 = upload_1.json()["import_id"]
    await client.post(f"/v1/admin/imports/{import_id_1}/apply", headers=_admin_headers(admin))
    await _apply_bound_to_test_session(db_session, monkeypatch, import_job_id=import_id_1)

    second_file = _build_xlsx(
        [
            [
                external_id, name, category, "Описание", sku_a,
                "цвет=красный", "999.00", "", "5", "", "",
            ],
            [
                external_id, name, category, "Описание", sku_b,
                "цвет=синий", "888.00", "", "3", "", "",
            ],
        ]
    )

    upload_2 = await client.post(
        "/v1/admin/imports/xlsx",
        headers=_admin_headers(admin),
        files={"file": ("second.xlsx", second_file, "application/octet-stream")},
    )
    assert upload_2.json()["preview"]["created_estimate"] == 0
    assert upload_2.json()["preview"]["updated_estimate"] == 2
    import_id_2 = upload_2.json()["import_id"]
    await client.post(f"/v1/admin/imports/{import_id_2}/apply", headers=_admin_headers(admin))
    await _apply_bound_to_test_session(db_session, monkeypatch, import_job_id=import_id_2)

    status_2 = await client.get(f"/v1/admin/imports/{import_id_2}", headers=_admin_headers(admin))
    assert status_2.json()["created_count"] == 0
    assert status_2.json()["updated_count"] == 2

    products_count = await db_session.scalar(
        select(func.count()).select_from(Product).where(Product.external_id == external_id)
    )
    assert products_count == 1

    variant_a = await db_session.scalar(select(ProductVariant).where(ProductVariant.sku == sku_a))
    variant_b = await db_session.scalar(select(ProductVariant).where(ProductVariant.sku == sku_b))
    assert variant_a is not None
    assert variant_a.price == Decimal("999.00")
    assert variant_a.stock_qty == 5
    assert variant_b is not None
    assert variant_b.price == Decimal("888.00")
    assert variant_b.stock_qty == 3


@pytest.mark.asyncio
async def test_malformed_rows_land_in_error_report_without_breaking_import(
    client: httpx.AsyncClient, db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    admin = await _make_admin(db_session)
    category = _slug("Категория")
    good_sku = _slug("sku-good")
    good_name = f"Хороший товар {good_sku}"

    file_bytes = _build_xlsx(
        [
            [None, good_name, category, "Описание", good_sku, "", "150.00", "", "10", "", ""],
            [None, "Без SKU и категории", "", "Описание", "", "", "150.00", "", "10", "", ""],
            [
                None, "Отрицательная цена текстом", category, "Описание", _slug("sku"),
                "", "abc", "", "10", "", "",
            ],
        ]
    )

    upload_response = await client.post(
        "/v1/admin/imports/xlsx",
        headers=_admin_headers(admin),
        files={"file": ("broken.xlsx", file_bytes, "application/octet-stream")},
    )
    import_id = upload_response.json()["import_id"]

    await client.post(f"/v1/admin/imports/{import_id}/apply", headers=_admin_headers(admin))
    await _apply_bound_to_test_session(db_session, monkeypatch, import_job_id=import_id)

    status_response = await client.get(
        f"/v1/admin/imports/{import_id}", headers=_admin_headers(admin)
    )
    body = status_response.json()
    assert body["status"] == "completed"
    assert body["created_count"] == 1
    assert body["error_count"] == 2
    assert body["errors_report_url"] is not None

    good_variant = await db_session.scalar(
        select(ProductVariant).where(ProductVariant.sku == good_sku)
    )
    assert good_variant is not None
    assert good_variant.price == Decimal("150.00")


@pytest.mark.asyncio
@pytest.mark.skipif(
    not s3_public_url_reachable(),
    reason="S3 public URL host is not reachable from this environment (e.g. inside a "
    "container, where 'localhost' isn't the host machine a browser would use)",
)
async def test_import_downloads_photos_into_s3(
    client: httpx.AsyncClient, db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    from app.config import settings
    from app.core.storage import ensure_bucket_exists, generate_presigned_url, get_s3_client

    ensure_bucket_exists()
    source_key = f"imports/test-source/{uuid.uuid4().hex}.jpg"
    get_s3_client().put_object(
        Bucket=settings.s3_bucket, Key=source_key, Body=b"fake-jpeg-bytes", ContentType="image/jpeg"
    )
    photo_url = generate_presigned_url(source_key)

    admin = await _make_admin(db_session)
    category = _slug("Категория")
    sku = _slug("sku")
    name = _slug("Товар с фото")

    file_bytes = _build_xlsx(
        [[None, name, category, "Описание", sku, "", "150.00", "", "10", "", photo_url]]
    )

    upload_response = await client.post(
        "/v1/admin/imports/xlsx",
        headers=_admin_headers(admin),
        files={"file": ("with-photo.xlsx", file_bytes, "application/octet-stream")},
    )
    import_id = upload_response.json()["import_id"]
    await client.post(f"/v1/admin/imports/{import_id}/apply", headers=_admin_headers(admin))
    await _apply_bound_to_test_session(db_session, monkeypatch, import_job_id=import_id)

    variant = await db_session.scalar(select(ProductVariant).where(ProductVariant.sku == sku))
    assert variant is not None
    product = await db_session.get(Product, variant.product_id)
    assert product is not None
    await db_session.refresh(product, attribute_names=["images"])
    assert len(product.images) == 1
